"""Check a monthly tab of the Customer Deposit workbook against the QBO 2115 GL, customer by customer.

For every customer:job it rolls last month's EOM forward with this month's 2115 activity
(credits = new deposits, debits = deposit released to revenue by the final invoice) and
compares the result to the EOM on this month's tab.

GL input is either:
  * the QBO GeneralLedger report JSON for account 2115 (what the n8n workflow pulls), or
  * an .xlsx/.csv export of the "CD Current Month GL" custom report
    (needs Transaction date, Name, Debit and Credit columns; the header row is found automatically).

Usage:
  python check_tab_against_gl.py WORKBOOK.xlsx --prev "Aug 26" --cur "Sep 26" --gl gl.json \
      --start 2026-09-01 --end 2026-09-25
"""
import argparse
import collections
import datetime as dt
import json
import re

import openpyxl

# Tab columns (0-based): A Customer, G DEPOSIT (= EOM), J BOM Dep, K Dep to Rev, L New Dep this Month.
COL_CUSTOMER, COL_Q, COL_EOM = 0, 1, 6


def key(name):
    return re.sub(r"\s+", " ", str(name or "")).strip().lower()


def read_tab(wb, tab):
    """EOM deposit per customer:job on a monthly tab (rows with a Q Number)."""
    eom, disp = collections.defaultdict(float), {}
    for r in wb[tab].iter_rows(min_row=2, values_only=True):
        if not r[COL_CUSTOMER] or r[COL_Q] is None:
            continue
        k = key(r[COL_CUSTOMER])
        disp[k] = r[COL_CUSTOMER]
        if isinstance(r[COL_EOM], (int, float)):
            eom[k] += r[COL_EOM]
    return eom, disp


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def num(v):
    if v in (None, ""):
        return 0.0
    return float(str(v).replace(",", "").replace("$", ""))


def read_gl_json(path):
    report = json.load(open(path))
    titles = [c["ColTitle"] for c in report["Columns"]["Column"]]
    i_date, i_name, i_dr, i_cr = (titles.index(t) for t in ("Date", "Name", "Debit", "Credit"))
    lines = []

    def walk(rows):
        for r in rows or []:
            if "Rows" in r:
                walk(r["Rows"].get("Row", []))
            if "ColData" not in r:
                continue
            c = [x.get("value") for x in r["ColData"]]
            if c[0] == "Beginning Balance":
                continue
            lines.append(dict(date=c[i_date], name=c[i_name], debit=num(c[i_dr]), credit=num(c[i_cr])))

    walk(report["Rows"]["Row"])
    return lines


def read_gl_table(path):
    if path.endswith(".csv"):
        import csv
        rows = list(csv.reader(open(path, newline="", encoding="utf-8-sig")))
    else:
        rows = list(openpyxl.load_workbook(path, data_only=True).active.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and "Name" in r and "Credit" in r)
    hdr = list(rows[hdr_i])
    ci = {h: hdr.index(h) for h in ("Transaction date", "Name", "Debit", "Credit")}
    lines = []
    for r in rows[hdr_i + 1:]:
        d = to_date(r[ci["Transaction date"]])
        if d and r[ci["Name"]]:
            lines.append(dict(date=d, name=r[ci["Name"]], debit=num(r[ci["Debit"]]), credit=num(r[ci["Credit"]])))
    return lines


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--prev", required=True, help='prior month tab, e.g. "Aug 26"')
    ap.add_argument("--cur", required=True, help='month being checked, e.g. "Sep 26"')
    ap.add_argument("--gl", required=True, help="2115 GL: QBO report .json, or custom report .xlsx/.csv")
    ap.add_argument("--start", required=True, help="first day of the month, YYYY-MM-DD")
    ap.add_argument("--end", required=True, help="cut-off date the tab was tied out at, YYYY-MM-DD")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.workbook, data_only=True)
    prev, disp = read_tab(wb, a.prev)
    cur, disp2 = read_tab(wb, a.cur)
    disp.update(disp2)
    gl = read_gl_json(a.gl) if a.gl.endswith(".json") else read_gl_table(a.gl)

    new, rel = collections.defaultdict(float), collections.defaultdict(float)
    for l in gl:
        if a.start <= l["date"] <= a.end:
            new[key(l["name"])] += l["credit"]
            rel[key(l["name"])] += l["debit"]
            disp.setdefault(key(l["name"]), l["name"])

    diffs = []
    for k in sorted(set(prev) | set(cur) | set(new) | set(rel)):
        expected = round(prev[k] + new[k] - rel[k], 2)
        if abs(cur[k] - expected) > 0.01:
            diffs.append((disp[k], prev[k], new[k], rel[k], expected, cur[k]))

    print(f"{a.cur}: BOM {sum(prev.values()):,.2f} + new {sum(new.values()):,.2f} "
          f"- dep to rev {sum(rel.values()):,.2f} = {sum(prev.values()) + sum(new.values()) - sum(rel.values()):,.2f}; "
          f"tab EOM {sum(cur.values()):,.2f}")
    if not diffs:
        print("Every customer:job ties to the GL.")
    for name, p, n, r, e, c in diffs:
        print(f"  {name[:60]:60} prev {p:>12,.2f} +new {n:>11,.2f} -rel {r:>11,.2f} = {e:>12,.2f}  tab {c:>12,.2f}  diff {c - e:>11,.2f}")


if __name__ == "__main__":
    main()
