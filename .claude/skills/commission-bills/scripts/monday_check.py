"""
Compare the monday.com Invoicing board's commissions for a payroll date with
the commission workbook.

    python3 monday_check.py <monday items.json> <YYYY-MM-DD> [--workbook book.xlsx] [--xlsx check.xlsx]

<monday items.json> is the result of the monday get_board_items_page call on
the Invoicing board (7469134121), filtered to items with any commission date
equal to the payroll date (see SKILL.md). It can be the saved tool result as
is; the "items" list is found inside it.

Without --workbook it lists what monday has for that date. With it, it
reports:
  * on monday for the date but not on the sheet (a missing commission);
  * on the sheet but not on monday for the date;
  * reps, rates or the sales-assistant 0.5% that differ between the two.
With --xlsx the findings are added as a "Monday check" tab to that workbook.
"""
import argparse
import json
import os
import re
import sys

WEBAPP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "..", "webapp")
sys.path.insert(0, WEBAPP)
import commission_bills as cb  # noqa: E402

# (label, Panda # column, commission date column, sheet stage)
TYPES = [
    ("Deposit", "numeric_mknr81a1", "date_mkme9g50", "Deposit"),
    ("Final", "numeric_mknrjqn3", "date_mkmengp1", "Final"),
    ("2nd Deposit", "numeric_mknrg083", "date_mkmeq670", "Deposit"),
    ("CO/Expedite", "numeric_mknrwgf7", "date_mkme5mt2", None),
    ("2nd CO/Expedite", "numeric_mknraf29", "date_mknrtqdn", None),
]
REPS = [("dropdown_mkkh6zyv", "text_mkm5bbak"), ("dup__of_in_house_rep_s__mkm0t75w", "text_mkm5bkxx"),
        ("dropdown_mkm0y8mx", "text_mkm5bv3q")]
SALES_ASSISTANT = ("dropdown_mkkrykjx", "text_mkm54qj1")
OUTSIDE_REP = ("outside_rep_name_mkm23atc", "text_mkm5dgey")
COLUMNS = (["text", "text_mkk9yngf", "text_mkm2tqcj"] + [c for t in TYPES for c in t[1:3]]
           + [c for pair in REPS + [SALES_ASSISTANT, OUTSIDE_REP] for c in pair])
NOT_A_PERSON = {"", "n/a", "na", "no", "none", "other"}


def find_items(obj):
    if isinstance(obj, dict):
        if isinstance(obj.get("items"), list):
            return obj["items"]
        for v in obj.values():
            got = find_items(v)
            if got is not None:
                return got
    if isinstance(obj, list):
        for v in obj:
            got = find_items(v)
            if got is not None:
                return got
    if isinstance(obj, str) and obj.lstrip().startswith("{"):
        try:
            return find_items(json.loads(obj))
        except ValueError:
            return None
    return None


def val(item, col):
    v = (item.get("column_values") or {}).get(col)
    if isinstance(v, dict):
        v = v.get("text") or ", ".join(map(str, v.get("ids") or [])) or None
    return str(v).strip() if v not in (None, "") else ""


def num(s):
    try:
        return float(str(s).replace(",", "").replace("$", ""))
    except ValueError:
        return None


def pct(s):
    """'4%', '3.5', '4%+1% = 5%', '0.5%' -> fraction; 'N/A' or blank -> None."""
    found = re.findall(r"\d+(?:\.\d+)?", str(s or "").replace("..", "."))
    return float(found[-1]) / 100 if found else None


def note_rate(notes, name):
    """A rep's rate written in the commission notes, e.g. 'Derrick 6 / Panda 3',
    '6% Derrick / 3% Panda', '9% - 4% Stefan / 5% Panda'. The notes win over the
    % column, which often holds the house split instead."""
    first = re.escape(name.split()[0]) if name else ""
    if not first or not notes:
        return None
    m = (re.search(rf"\b{first}\b\s*[:=-]?\s*(\d+(?:\.\d+)?)\s*%?", notes, re.I)
         or re.search(rf"(\d+(?:\.\d+)?)\s*%?\s*-?\s*{first}\b", notes, re.I))
    return float(m.group(1)) / 100 if m else None


def person(s):
    return "" if cb.norm(s) in NOT_A_PERSON else s.strip()


def same_person(monday_name, vendor_name):
    a, b = set(cb.norm(monday_name).split()), set(cb.norm(vendor_name).split())
    return bool(a) and (a <= b or b <= a)


def monday_lines(items, date):
    """One entry per (item, commission type) dated on the payroll date."""
    out = []
    for it in items:
        for label, panda_col, date_col, stage in TYPES:
            if val(it, date_col) != date:
                continue
            notes = val(it, "text_mkm2tqcj")
            reps = [(person(val(it, n)), note_rate(notes, person(val(it, n))) or pct(val(it, p)))
                    for n, p in REPS if person(val(it, n))]
            sa = (person(val(it, SALES_ASSISTANT[0])), pct(val(it, SALES_ASSISTANT[1])))
            outside = (person(val(it, OUTSIDE_REP[0])), pct(val(it, OUTSIDE_REP[1])))
            out.append({
                "item": it.get("name", ""), "url": it.get("url", ""), "customer": val(it, "text"),
                "q": val(it, "text_mkk9yngf"), "type": label, "stage": stage,
                "panda": num(val(it, panda_col)), "reps": reps,
                "sa": sa if sa[0] and sa[1] else None, "outside": outside if outside[0] else None,
                "notes": notes,
            })
    return out


def describe(m):
    reps = " + ".join(f"{n} {r * 100:g}%" if r else f"{n} (no %)" for n, r in m["reps"]) or "no in-house rep"
    if m["sa"]:
        reps += f"; SA {m['sa'][0]} {m['sa'][1] * 100:g}%"
    if m["outside"]:
        reps += f"; outside rep {m['outside'][0]}" + (f" {m['outside'][1] * 100:g}%" if m["outside"][1] else "")
    return reps


def compare(mlines, rows, vendor_map, second_vendor):
    findings = []
    used = set()
    for m in mlines:
        cands = [i for i, r in enumerate(rows) if r["q_number"] == m["q"] and i not in used]
        exact = [i for i in cands if m["panda"] is not None and r_panda(rows[i]) is not None
                 and abs(r_panda(rows[i]) - m["panda"]) <= 1]
        pick = (exact or cands or [None])[0]
        if pick is None:
            amt = ""
            if m["panda"] and m["reps"]:
                amt = " -> " + ", ".join(f"{n} ${m['panda'] * r:,.2f}" for n, r in m["reps"] if r)
            findings.append(("NOT ON SHEET", m, None, f"{m['type']} P# {m['panda'] or '?':,} {describe(m)}{amt}"))
            continue
        used.add(pick)
        r = rows[pick]
        issues = []
        if not exact:
            issues.append(f"Panda # sheet {r_panda(r):,.0f} vs monday {m['panda'] or 0:,.0f}")
        if m["stage"] and m["stage"] not in r["doc_number"]:
            issues.append(f"sheet says {r['doc_number']}, monday says {m['type']}")
        mapped = vendor_map.get(cb.norm(r["rep"]))
        vendors = [mapped] if isinstance(mapped, str) else [v if isinstance(v, str) else v["vendor"] for v in (mapped or [])]
        mreps = [n for n, _ in m["reps"]]
        missing = [n for n in mreps if not any(same_person(n, v) for v in vendors)]
        extra = [v for v in vendors if not any(same_person(n, v) for n in mreps)]
        if missing or extra:
            issues.append(f"reps: sheet {r['rep']} ({' + '.join(vendors) or 'unmapped'}) vs monday {' + '.join(mreps) or 'none'}")
        rates = {round(rr, 5) for _, rr in m["reps"] if rr}
        if r["rate"] is not None and rates and round(r["rate"], 5) not in rates:
            issues.append(f"rate: sheet {r['rate'] * 100:g}% vs monday {' / '.join(f'{x * 100:g}%' for x in sorted(rates))}")
        if m["sa"] and not r["second"]:
            issues.append(f"monday has SA {m['sa'][0]} {m['sa'][1] * 100:g}% but sheet has no 2nd Comm")
        elif m["sa"] and second_vendor and not same_person(m["sa"][0], second_vendor):
            issues.append(f"2nd Comm goes to {second_vendor} but monday SA is {m['sa'][0]}")
        elif r["second"] and not m["sa"]:
            issues.append("sheet has 2nd Comm but monday has no sales-assistant %")
        if issues:
            findings.append(("DIFFERS", m, r, "; ".join(issues)))
    for i, r in enumerate(rows):
        if i not in used:
            findings.append(("NOT ON MONDAY FOR THIS DATE", None, r,
                             f"{r['doc_number']} {r['rep']} ${r['amount']:,.2f} (check its commission date on monday)"))
    return findings


def r_panda(r):
    return r.get("panda")


def write_tab(path, date, findings, mlines):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.load_workbook(path) if os.path.exists(path) else openpyxl.Workbook()
    if "Monday check" in wb.sheetnames:
        del wb["Monday check"]
    ws = wb.create_sheet("Monday check")
    ws.append([f"monday.com Invoicing board vs sheet, commission date {date}"])
    ws["A1"].font = Font(bold=True)
    ws.append([f"{len(mlines)} commission(s) dated {date} on monday; {len(findings)} finding(s)"])
    ws.append([])
    hdr = ["Finding", "Q#", "Job", "Customer", "monday type", "Panda #", "monday reps / %", "Sheet bill", "Detail", "monday link"]
    ws.append(hdr)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F9D45")
    colors = {"NOT ON SHEET": "F8CBAD", "DIFFERS": "FFE699", "NOT ON MONDAY FOR THIS DATE": "DDEBF7"}
    for kind, m, r, detail in findings:
        ws.append([kind, (m or {}).get("q") or (r or {}).get("q_number"), (m or {}).get("item") or (r or {}).get("customer"),
                   (m or {}).get("customer", ""), (m or {}).get("type", ""), (m or {}).get("panda") or (r or {}).get("panda"),
                   describe(m) if m else "", (r or {}).get("doc_number", ""), detail, (m or {}).get("url", "")])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=colors[kind])
    for col, w in zip("ABCDEFGHIJ", [26, 12, 30, 28, 14, 11, 48, 15, 70, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5):
        row[8].alignment = Alignment(wrap_text=True, vertical="top")
        row[5].number_format = "#,##0"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("items_json")
    ap.add_argument("date", help="payroll / commission date, YYYY-MM-DD")
    ap.add_argument("--workbook")
    ap.add_argument("--xlsx", help="add a 'Monday check' tab to this workbook (created if missing)")
    args = ap.parse_args()

    items = find_items(json.load(open(args.items_json)))
    if items is None:
        raise SystemExit("No 'items' list found in " + args.items_json)
    mlines = monday_lines(items, args.date)
    print(f"monday: {len(mlines)} commission(s) dated {args.date} on {len(items)} item(s)")
    if not args.workbook:
        for m in mlines:
            print(f"  {m['q']:<12} {m['type']:<12} P# {m['panda'] or 0:>10,.0f}  {m['item'][:35]:<35} {describe(m)}")
        return 0

    _, rows = cb.read_rows(args.workbook)
    vmap_raw = json.load(open(os.path.join(WEBAPP, "commission_vendors.json")))
    vmap = {cb.norm(k): v for k, v in vmap_raw.items() if not k.startswith("_")}
    findings = compare(mlines, rows, vmap, vmap.get(cb.norm(cb.SECOND_COMM_KEY)))
    order = {"NOT ON SHEET": 0, "DIFFERS": 1, "NOT ON MONDAY FOR THIS DATE": 2}
    findings.sort(key=lambda f: order[f[0]])
    for kind, m, r, detail in findings:
        label = (m["q"] + " " + m["item"]) if m else (r["q_number"] + " " + r["customer"])
        print(f"[{kind}] {label[:55]}: {detail}")
    if not findings:
        print("Sheet matches monday for this date.")
    if args.xlsx:
        write_tab(args.xlsx, args.date, findings, mlines)
        print(f"Added 'Monday check' tab to {args.xlsx}")
    return 1 if findings else 0


if __name__ == "__main__":
    sys.exit(main())
